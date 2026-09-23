"""叙事文本抽取：真实文档 -> 可直接喂 src/ragspine/retrieval/chunking/chunking.py 的「文档文本 + 定位段」。

叙事通路（Narrative Path）吃的是归因 / 监管 / 会议纪要类文本（docs/architecture.md「Channels」），
本模块只做确定性文本抽取，零 OCR、零 LLM：

    - pptx：python-pptx 遍历每页文本框（含占位符）+ 演讲者备注，按 slide 聚合；
      表格内容跳过（表格数字归结构化通路管）。定位串 'slide={N},frame={M}'
      （M 只给产出非空文本的非表格文本框编号，1-based）/ 'slide={N},notes'。
    - PDF（数字型）：pypdfium2 textpage 逐页取文本，定位串 'page={N}'（真实页号）。
      扫描型页（无文本层）本期不做（OCR 线另有归属）：跳过、计数并告警。
    - docx（W3b）：docspine 遍历正文按文档顺序取段落文本；表格内容跳过（表格数字归
      结构化通路管，与 pptx 同口径）。定位串 'para={N}'（N 只给产出非空文本的段落编号，
      1-based）。docspine 是可选 [doc] extra，惰性 import。
    - txt（纯文本）：按空行拆成段落块，逐块归一化后成段；定位串 'para={N}'（N 只给
      产出非空文本的块编号，1-based，与 docx 同口径）。零三方依赖。
    - md（DI markdown，Azure Document Intelligence 风格）：extraction.di_markdown 解析成
      页 → 块，按（页, 标题路径）连续分组成段；定位串 'page={N}'（N=物理页序 DiPage.index，
      与 PDF 同口径；PageNumber 标签不参与）。表格线性化为「行标题 | 列标题: 值」行、
      图 caption 在前；页眉/页脚/页码注释不进正文。无 PageBreak 的普通 markdown 整份 = 第 1 页。

统一返回 NarrativeDoc：segments 各段文本内部以 '\\n' 分行、to_text() 以空行
（'\\n\\n'）连接各段，正好配合 chunking 的「非空行 = 段落」切分契约。
"""

from dataclasses import dataclass, field
from pathlib import Path

import pypdfium2 as pdfium
from pptx import Presentation

from ragspine.extraction.di_markdown.models import Figure, Table, TableGrid
from ragspine.extraction.di_markdown.parse import parse_di_markdown
from ragspine.extraction.extractors.pptx_styled_extractor import compute_file_hash

# 本期支持的叙事来源后缀（扫描件 / 其它格式归别的线）。
SUPPORTED_SUFFIXES = {".pptx", ".pdf", ".docx", ".docm", ".txt", ".md"}


@dataclass
class NarrativeSegment:
    """一个叙事文本段：文本 + 原文定位串。

    source_locator 取值约定：
        'slide={N},frame={M}'  pptx 第 N 页第 M 个产出文本的文本框（均 1-based）。
        'slide={N},notes'      pptx 第 N 页演讲者备注。
        'page={N}'             PDF 第 N 页（1-based，真实页号，跳过页不占用）。
        'para={N}'             docx 第 N 个产出非空文本的段落（1-based，空段不占用）。

    heading_path: 所属标题路径（DI markdown 填；其余抽取器留空）。
    """

    text: str
    source_locator: str
    heading_path: tuple[str, ...] = ()


@dataclass
class NarrativeDoc:
    """一个文档的叙事抽取结果（喂 chunking 的统一中间表示）。

    字段语义约定：
        doc_id:        源文件名（与全库血缘约定一致）。
        file_hash:     源文件内容 sha256 十六进制串（版本血缘）。
        segments:      叙事文本段列表（按 slide/页序）。
        skipped_pages: 无文本层被跳过的 PDF 页数（pptx 恒为 0）。
        warnings:      告警汇聚（如逐页跳过原因）。
    """

    doc_id: str
    file_hash: str
    segments: list[NarrativeSegment] = field(default_factory=list)
    skipped_pages: int = 0
    warnings: list[str] = field(default_factory=list)

    def to_text(self) -> str:
        """各段以空行连接成文档级纯文本（chunking 的输入契约）。"""
        return "\n\n".join(seg.text for seg in self.segments)


def _clean_block(text: str) -> str:
    """块内文本归一化：逐行折叠内部空白、丢空行，行间以 '\\n' 连接。

    pptx 软换行（'\\v'）视同换行；返回空串表示该块无实质文本。
    """
    lines = (" ".join(line.split()) for line in text.replace("\v", "\n").splitlines())
    return "\n".join(line for line in lines if line)


def extract_pptx_narrative(path: str | Path) -> NarrativeDoc:
    """抽取一个 pptx 的叙事文本：每页文本框（含占位符）+ 演讲者备注。

    - 表格 shape 跳过（表格数字归结构化通路）。
    - 仅产出非空文本的文本框获得 frame 序号（与 pptx_styled_extractor 的
      box_no 约定一致）；备注段排在该页文本框段之后。
    """
    path = Path(path)
    doc = NarrativeDoc(doc_id=path.name, file_hash=compute_file_hash(path))
    prs = Presentation(str(path))

    for slide_no, slide in enumerate(prs.slides, start=1):
        frame_no = 0
        for shape in slide.shapes:
            if shape.has_table or not shape.has_text_frame:
                continue
            text = _clean_block(shape.text_frame.text)
            if not text:
                continue
            frame_no += 1
            doc.segments.append(
                NarrativeSegment(
                    text=text,
                    source_locator=f"slide={slide_no},frame={frame_no}",
                )
            )
        if slide.has_notes_slide:
            notes = _clean_block(slide.notes_slide.notes_text_frame.text)
            if notes:
                doc.segments.append(
                    NarrativeSegment(
                        text=notes,
                        source_locator=f"slide={slide_no},notes",
                    )
                )
    return doc


def extract_pdf_narrative(path: str | Path) -> NarrativeDoc:
    """抽取一个数字型 PDF 的叙事文本：pypdfium2 textpage 逐页取文本按页聚合。

    无文本层的页（扫描型形态）跳过：skipped_pages 计数 + warnings 点名页号，
    OCR 归扫描线处理，本模块绝不臆造内容。
    """
    path = Path(path)
    doc = NarrativeDoc(doc_id=path.name, file_hash=compute_file_hash(path))
    pdf = pdfium.PdfDocument(str(path))
    try:
        for idx in range(len(pdf)):
            page = pdf[idx]
            try:
                textpage = page.get_textpage()
                try:
                    raw = textpage.get_text_range()
                finally:
                    textpage.close()
            finally:
                page.close()
            text = _clean_block(raw)
            if text:
                doc.segments.append(
                    NarrativeSegment(
                        text=text,
                        source_locator=f"page={idx + 1}",
                    )
                )
            else:
                doc.skipped_pages += 1
                doc.warnings.append(f"page={idx + 1}: 无文本层，跳过（扫描页归 OCR 线处理）")
    finally:
        pdf.close()
    return doc


def extract_docx_narrative(path: str | Path) -> NarrativeDoc:
    """抽取一个 .docx 的叙事文本：正文段落按文档顺序聚合（W3b）。

    - 表格 block 跳过（表格数字归结构化通路，与 pptx 同口径）。
    - 仅产出非空文本的段落获得 para 序号；locator='para={N}'（N=非空段序，1-based）。

    docspine 是可选 [doc] extra，惰性 import（未装 [doc] 也能 import 本模块、跑 pptx/pdf）。
    """
    path = Path(path)
    doc = NarrativeDoc(doc_id=path.name, file_hash=compute_file_hash(path))

    import docspine

    parsed = docspine.open(str(path))
    para_no = 0
    for block in parsed.body():
        if block["kind"] != "paragraph":
            continue
        text = _clean_block(block["text"])
        if not text:
            continue
        para_no += 1
        doc.segments.append(
            NarrativeSegment(
                text=text,
                source_locator=f"para={para_no}",
            )
        )
    return doc


def extract_txt_narrative(path: str | Path) -> NarrativeDoc:
    """抽取一个纯文本 .txt 的叙事文本：按空行拆段落块，逐块归一化聚合。

    - UTF-8 读取，容错 errors='replace'（避免个别坏字节中断整批，与 no-fabrication 无冲突）。
    - 仅产出非空文本的块获得 para 序号；locator='para={N}'（N=非空块序，1-based，
      空块不占用，与 docx 同口径）。零三方依赖、确定性。
    """
    path = Path(path)
    doc = NarrativeDoc(doc_id=path.name, file_hash=compute_file_hash(path))

    raw = path.read_text(encoding="utf-8", errors="replace")
    para_no = 0
    for block in raw.split("\n\n"):
        text = _clean_block(block)
        if not text:
            continue
        para_no += 1
        doc.segments.append(
            NarrativeSegment(
                text=text,
                source_locator=f"para={para_no}",
            )
        )
    return doc


def _dedupe(parts: list[str]) -> list[str]:
    """去空、按首次出现去重（保序）。"""
    out: list[str] = []
    for part in parts:
        if part and part not in out:
            out.append(part)
    return out


def _row_values(grid: TableGrid, r: int, start: int) -> list[tuple[int, str]]:
    """第 r 行自 start 列起的非空值 (列, 文本)；横向合并只在锚点列出现一次。"""
    out: list[tuple[int, str]] = []
    for c in range(start, grid.n_cols):
        anchor = grid.anchor_at(r, c)
        if anchor is not None and anchor.text and (anchor.col == c or c == start):
            out.append((c, anchor.text))
    return out


def _linearize_table(grid: TableGrid) -> list[str]:
    """表格 → 文本行：caption 在前；有表头时每个值一行「行标题 | 列标题: 值」
    （多行表头按列 ' / ' 拼接去重，左上角表头格单独一行作表格语境），无表头时整行 ' | ' 连接。
    """
    lines = [grid.caption] if grid.caption else []
    rows = grid.rows
    n_header = grid.header_row_count
    if n_header == 0 or n_header >= grid.n_rows or grid.n_cols < 2:
        for r in range(grid.n_rows):
            line = " | ".join(text for _, text in _row_values(grid, r, 0))
            if line:
                lines.append(line)
        return lines
    col_headers = [
        " / ".join(_dedupe([rows[r][c] for r in range(n_header)])) for c in range(grid.n_cols)
    ]
    if col_headers[0]:
        lines.append(col_headers[0])
    for r in range(n_header, grid.n_rows):
        row_title = rows[r][0]
        values = []
        for c, value in _row_values(grid, r, 1):
            label = f"{col_headers[c]}: {value}" if col_headers[c] else value
            values.append(f"{row_title} | {label}" if row_title else label)
        lines.extend(values or ([row_title] if row_title else []))
    return lines


def extract_di_markdown_narrative(path: str | Path) -> NarrativeDoc:
    """抽取一个 DI markdown（.md）的叙事文本：每页内按标题路径连续分组，一组一段。

    locator='page={N}'（N=物理页序），段带 heading_path；标题文本作为段首行保留。
    UTF-8 读取，容错 errors='replace'。纯 stdlib、确定性、不涉及任何公司。
    """
    path = Path(path)
    doc = NarrativeDoc(doc_id=path.name, file_hash=compute_file_hash(path))
    parsed = parse_di_markdown(path.read_text(encoding="utf-8", errors="replace"))

    for page in parsed.pages:
        groups: list[tuple[tuple[str, ...], list[str]]] = []
        for block in page.blocks:
            if isinstance(block, Table):
                lines = _linearize_table(block.grid)
            elif isinstance(block, Figure):
                lines = [block.caption or "", block.text]
            else:  # Heading / Paragraph
                lines = [block.text]
            if not groups or groups[-1][0] != block.heading_path:
                groups.append((block.heading_path, []))
            groups[-1][1].extend(lines)
        for heading_path, lines in groups:
            text = _clean_block("\n".join(lines))
            if text:
                doc.segments.append(
                    NarrativeSegment(
                        text=text,
                        source_locator=f"page={page.index}",
                        heading_path=heading_path,
                    )
                )
    return doc


def extract_narrative(path: str | Path) -> NarrativeDoc:
    """按后缀分发到对应抽取器；不支持的后缀 ValueError。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return extract_pptx_narrative(path)
    if suffix == ".pdf":
        return extract_pdf_narrative(path)
    if suffix in (".docx", ".docm"):
        return extract_docx_narrative(path)
    if suffix == ".txt":
        return extract_txt_narrative(path)
    if suffix == ".md":
        return extract_di_markdown_narrative(path)
    raise ValueError(f"不支持的叙事来源类型：{path.name}（仅 .pptx / .pdf / .docx / .txt / .md）")
