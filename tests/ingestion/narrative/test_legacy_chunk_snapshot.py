"""叙事入库旧路径字节不变守护：pptx / pdf / docx / txt 在「按段切块」开关关闭时，
块库内容（除入库时间戳外全部字段）与冻结快照逐字节一致。

快照摘要取自引入 DI markdown / 按段切块开关之前的实现（HEAD d6b11e0）；
开关默认值与显式 False 两种调用都必须命中同一摘要。
"""

import hashlib
import json
import os
from dataclasses import asdict

import pytest
import rootutils

ROOT_DIR = rootutils.setup_root(os.getcwd(), indicator=".project-root", pythonpath=True)

from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas

from ragspine.ingestion.narrative.narrative_ingest import ingest_narrative
from ragspine.retrieval.chunking.chunk_store import ChunkStore
from ragspine.retrieval.chunking.chunker import make_chunker

pytest.importorskip("docspine", reason="docspine 未安装（[doc]）")

_LONG = "Hong Kong REVENUE grew as agency productivity improved. " * 12

# 冻结快照：{chunker spec: sha256}。改动旧路径导致摘要变化 = 行为回归。
_FROZEN = {
    "none": "ac1b8b36caed887f6ea1ee2b1828fbd7f7c2268674daf9d927be8a1626237d76",
    "parent_child": "316359a978cfcc59b044497a53c9ea1d74ceb3a0324d166dc294f06fca85fd9b",
}


def _make_corpus(tmp_path, make_docx) -> list:
    deck = tmp_path / "deck_FY2024.pptx"
    prs = Presentation()
    for texts, notes in ((["Slide one title", _LONG], "Speaker note one."), (["Slide two"], None)):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, text in enumerate(texts):
            tb = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5 + i * 0.9), Inches(8), Inches(0.8)
            )
            tb.text_frame.text = text
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(deck))

    pdf = tmp_path / "report_2025H1.pdf"
    c = Canvas(str(pdf), pagesize=(595, 842))
    for line in ("Page one narrative about claims.", None, "Page three outlook text."):
        if line is None:
            c.rect(40, 40, 500, 760, fill=1, stroke=0)
        else:
            c.setFont("Helvetica", 12)
            c.drawString(60, 700, line)
        c.showPage()
    c.save()

    docx = tmp_path / "memo.docx"
    make_docx(
        docx,
        [
            ("para", "Memo heading"),
            ("table", [["x", "1"]]),
            ("para", _LONG),
        ],
    )

    txt = tmp_path / "notes.txt"
    txt.write_text(f"First block.\n\n{_LONG}\n\nThird block.", encoding="utf-8")
    return [deck, pdf, docx, txt]


def _digest(store: ChunkStore) -> str:
    rows = []
    for c in store.iter_chunks():
        d = asdict(c)
        d.pop("ingested_at")
        rows.append(d)
    blob = json.dumps(rows, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


@pytest.mark.parametrize("spec", sorted(_FROZEN))
@pytest.mark.parametrize(
    "switch_kwargs", [{}, {"segment_chunking": False}], ids=["default", "explicit_off"]
)
def test_legacy_suffixes_byte_identical_when_switch_off(tmp_path, make_docx, spec, switch_kwargs):
    files = _make_corpus(tmp_path, make_docx)
    store = ChunkStore(tmp_path / "chunks.db")
    store.init_schema()
    try:
        report = ingest_narrative(files, store, chunker=make_chunker(spec), **switch_kwargs)
        assert [f.status for f in report.files] == ["ingested"] * 4
        assert _digest(store) == _FROZEN[spec]
    finally:
        store.close()
